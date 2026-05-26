from pathlib import Path
import math

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
TEMPLATE = Path(r"e:\学习\数字孪生大创\软著申请\第二次申请\基于数字孪生的实木加工生产线数据感知软件说明书.docx")
OUT_DOCX = ROOT / "双线性Hall高精度位置检测与补偿软件说明书.docx"
FIG_DIR = ROOT / "software_manual_figures"


def font(size=26, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


F_TITLE = font(32, True)
F_H = font(24, True)
F = font(20)
F_SMALL = font(17)


def draw_round(draw, xy, r, fill, outline="#8bbdd8", width=2):
    draw.rounded_rectangle(xy, radius=r, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, fnt, fill="#15384f"):
    lines = wrap_text(draw, text, fnt, box[2] - box[0] - 20)
    total_h = sum(draw.textbbox((0, 0), ln, font=fnt)[3] - draw.textbbox((0, 0), ln, font=fnt)[1] for ln in lines) + (len(lines) - 1) * 6
    y = box[1] + (box[3] - box[1] - total_h) / 2
    for ln in lines:
        bb = draw.textbbox((0, 0), ln, font=fnt)
        x = box[0] + (box[2] - box[0] - (bb[2] - bb[0])) / 2
        draw.text((x, y), ln, font=fnt, fill=fill)
        y += bb[3] - bb[1] + 6


def wrap_text(draw, text, fnt, max_w):
    lines = []
    current = ""
    for ch in text:
        test = current + ch
        if draw.textlength(test, font=fnt) <= max_w:
            current = test
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def arrow(draw, p1, p2, color="#466174", width=4):
    draw.line([p1, p2], fill=color, width=width)
    a = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    head = 14
    pts = [
        p2,
        (p2[0] - head * math.cos(a - 0.45), p2[1] - head * math.sin(a - 0.45)),
        (p2[0] - head * math.cos(a + 0.45), p2[1] - head * math.sin(a + 0.45)),
    ]
    draw.polygon(pts, fill=color)


def save_architecture():
    img = Image.new("RGB", (1400, 780), "#f7fbff")
    d = ImageDraw.Draw(img)
    d.text((40, 35), "双线性 Hall 高精度位置检测与补偿软件功能结构", font=F_TITLE, fill="#12384f")
    modules = [
        ("参数配置", "极对数、安装角、幅值、零偏、相位、ADC、噪声、采样率"),
        ("信号生成", "生成双线性 Hall 正余弦信号与磁场角、机械角"),
        ("误差注入", "零偏、幅值不一致、相位不正交、谐波、白噪声、量化"),
        ("补偿算法", "零偏/幅值/相位、谐波、LPF、α-β、离散 ESO"),
        ("工况仿真", "恒速、变速、冲击、负载扰动"),
        ("结果显示", "电压波形、相图、FFT、误差曲线、指标表"),
        ("报告导出", "PNG、CSV、仿真参数 JSON"),
    ]
    x0, y0, w, h = 55, 120, 300, 115
    for i, (name, desc) in enumerate(modules):
        row, col = divmod(i, 3)
        x = x0 + col * 430
        y = y0 + row * 185
        draw_round(d, (x, y, x + w, y + h), 18, "#ffffff", "#9ad2f2", 3)
        d.text((x + 18, y + 16), name, font=F_H, fill="#1672ad")
        for j, ln in enumerate(wrap_text(d, desc, F_SMALL, w - 34)):
            d.text((x + 18, y + 56 + j * 24), ln, font=F_SMALL, fill="#60788a")
        if i < len(modules) - 1:
            nx = x0 + ((i + 1) % 3) * 430
            ny = y0 + ((i + 1) // 3) * 185
            if col < 2 and row == (i + 1) // 3:
                arrow(d, (x + w + 18, y + h // 2), (nx - 18, ny + h // 2))
            elif col == 2:
                arrow(d, (x + w // 2, y + h + 16), (nx + w // 2, ny - 18))
    path = FIG_DIR / "fig1_software_architecture.png"
    img.save(path)
    return path


def save_workflow():
    img = Image.new("RGB", (1400, 650), "#f7fbff")
    d = ImageDraw.Draw(img)
    d.text((40, 35), "信号处理与补偿流程", font=F_TITLE, fill="#12384f")
    steps = [
        "两路 Hall 信号",
        "零偏估计",
        "幅值归一化",
        "相位/安装角补偿",
        "谐波识别与重构",
        "atan2 解角",
        "α-β / ESO 动态估计",
        "误差评价与导出",
    ]
    y = 190
    for i, s in enumerate(steps):
        x = 45 + i * 165
        draw_round(d, (x, y, x + 135, y + 82), 16, "#ffffff", "#86c8ee", 3)
        center_text(d, (x + 6, y + 5, x + 129, y + 77), s, F_SMALL)
        if i < len(steps) - 1:
            arrow(d, (x + 140, y + 41), (x + 160, y + 41))
    formulas = [
        "归一化：x=(H-O)/A",
        "相位补偿：xc'=(xc+xs·sinΔφ)/cosΔφ",
        "谐波补偿：e(θ)=Σ(Ak sin kθ+Bk cos kθ)",
        "解角：θm=atan2(xs,xc')/p",
    ]
    for i, txt in enumerate(formulas):
        draw_round(d, (90 + i * 310, 390, 360 + i * 310, 475), 12, "#eaf7ff", "#c4e5f7", 2)
        center_text(d, (100 + i * 310, 400, 350 + i * 310, 465), txt, F_SMALL)
    path = FIG_DIR / "fig2_signal_workflow.png"
    img.save(path)
    return path


def save_installation():
    img = Image.new("RGB", (1200, 720), "#f7fbff")
    d = ImageDraw.Draw(img)
    d.text((40, 35), "中空磁钢与双线性 Hall 安装示意", font=F_TITLE, fill="#12384f")
    cx, cy = 450, 370
    outer, inner = 190, 86
    d.ellipse((cx - outer, cy - outer, cx + outer, cy + outer), fill="#eaf7ff", outline="#7fbfe5", width=4)
    d.ellipse((cx - inner, cy - inner, cx + inner, cy + inner), fill="#ffffff", outline="#9fcde6", width=3)
    pole_count = 12
    for i in range(pole_count):
        a = -math.pi / 2 + i * 2 * math.pi / pole_count
        x = cx + math.cos(a) * 140
        y = cy + math.sin(a) * 140
        color = "#6ec6e8" if i % 2 else "#f49b69"
        d.ellipse((x - 22, y - 22, x + 22, y + 22), fill=color, outline="#ffffff", width=2)
        d.text((x - 6, y - 12), "N" if i % 2 else "S", font=F_SMALL, fill="#ffffff")
    hall_s = (cx, cy - outer - 56)
    hall_c = (cx + outer + 56, cy)
    draw_round(d, (hall_s[0] - 55, hall_s[1] - 22, hall_s[0] + 55, hall_s[1] + 22), 10, "#fff", "#d9812b", 3)
    center_text(d, (hall_s[0] - 55, hall_s[1] - 22, hall_s[0] + 55, hall_s[1] + 22), "Hall S", F_SMALL, "#d9812b")
    draw_round(d, (hall_c[0] - 55, hall_c[1] - 22, hall_c[0] + 55, hall_c[1] + 22), 10, "#fff", "#2ba86a", 3)
    center_text(d, (hall_c[0] - 55, hall_c[1] - 22, hall_c[0] + 55, hall_c[1] + 22), "Hall C", F_SMALL, "#2ba86a")
    d.arc((cx - 90, cy - 90, cx + 90, cy + 90), -90, 0, fill="#1672ad", width=4)
    d.text((cx + 20, cy - 88), "安装夹角约 90°", font=F_SMALL, fill="#1672ad")
    d.line((cx - outer - 80, cy, cx + outer + 80, cy), fill="#315b73", width=4)
    arrow(d, (cx, cy), (cx + 125, cy - 95), "#1672ad", 4)
    d.text((cx + 95, cy - 128), "磁场角 θmag", font=F_SMALL, fill="#1672ad")
    notes = [
        "1. 中空磁钢固定于转子或关节轴端，形成周期磁场。",
        "2. 两只线性 Hall 近似正交安装，输出 Hs/Hc。",
        "3. 安装角偏离 90° 会表现为相位不正交误差。",
        "4. 软件可调安装角、极对数、ADC 位数和误差强度。"
    ]
    draw_round(d, (760, 170, 1130, 500), 16, "#ffffff", "#bde0f5", 3)
    d.text((785, 195), "安装建模要点", font=F_H, fill="#1672ad")
    for i, n in enumerate(notes):
        for j, ln in enumerate(wrap_text(d, n, F_SMALL, 315)):
            d.text((785, 245 + i * 60 + j * 24), ln, font=F_SMALL, fill="#60788a")
    path = FIG_DIR / "fig3_hall_installation.png"
    img.save(path)
    return path


def save_ui_layout():
    img = Image.new("RGB", (1400, 720), "#f7fbff")
    d = ImageDraw.Draw(img)
    d.text((40, 35), "软件界面布局示意", font=F_TITLE, fill="#12384f")
    draw_round(d, (55, 105, 390, 650), 18, "#ffffff", "#9ad2f2", 3)
    d.text((85, 135), "左侧参数配置", font=F_H, fill="#1672ad")
    left = ["Hall/ADC参数", "误差注入开关", "补偿算法参数", "工况仿真参数", "一键运行与导出"]
    for i, item in enumerate(left):
        draw_round(d, (90, 190 + i * 76, 355, 240 + i * 76), 10, "#eaf7ff", "#c4e5f7", 2)
        center_text(d, (100, 194 + i * 76, 345, 236 + i * 76), item, F_SMALL)
    draw_round(d, (450, 105, 1340, 650), 18, "#ffffff", "#9ad2f2", 3)
    d.text((485, 135), "右侧结果工作区", font=F_H, fill="#1672ad")
    blocks = [
        ("动态安装示意", (490, 190, 840, 330)),
        ("电压波形 / 相图", (880, 190, 1270, 330)),
        ("FFT / 解角误差", (490, 370, 840, 510)),
        ("指标表 / 雷达图 / 导出", (880, 370, 1270, 510)),
    ]
    for title, box in blocks:
        draw_round(d, box, 12, "#f5fbff", "#c4e5f7", 2)
        center_text(d, box, title, F_SMALL)
    path = FIG_DIR / "fig4_ui_layout.png"
    img.save(path)
    return path


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", style=None, align=None, size=11, bold=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 14, bold=True, color=(22, 114, 173) if level == 1 else None)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def clear_document(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def add_figure(doc, path, caption, width_cm=14.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cap.add_run(caption)
    set_run_font(rr, size=10, bold=False, color=(96, 120, 138))


def build_doc():
    FIG_DIR.mkdir(exist_ok=True)
    figs = [save_architecture(), save_workflow(), save_installation(), save_ui_layout()]

    doc = Document(str(TEMPLATE)) if TEMPLATE.exists() else Document()
    clear_document(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("双线性 Hall 高精度位置检测与补偿软件说明书")
    set_run_font(run, size=24, bold=True)
    add_para(doc, "V1.0", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    for _ in range(8):
        add_para(doc)
    add_para(doc, "北京航空航天大学", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    add_para(doc, "2026-05-25", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    doc.add_page_break()

    add_heading(doc, "声明", 1)
    add_para(doc, "本手册由项目组编写，文档内容用于“双线性 Hall 高精度位置检测与补偿软件”的软件著作权申请、方案说明和内部交流。未经许可，不得以任何方式复制、传播或用于与本项目无关的用途。", size=11)
    for _ in range(3):
        add_para(doc)
    info = [
        "编 写 人：项目组",
        "使用部门：多轴手机云台项目组",
        "使 用 人：双线性 Hall 位置检测与补偿算法开发人员",
        "密    级：内部文档",
        "日    期：2026-05-25",
    ]
    for line in info:
        add_para(doc, line, size=11)

    doc.add_page_break()
    add_heading(doc, "1 引言", 1)
    add_heading(doc, "编写目的", 2)
    add_para(doc, "本说明书用于描述“双线性 Hall 高精度位置检测与补偿软件”的设计目标、运行环境、功能组成、算法流程和可行性。软件面向微型无刷电机、云台关节和大角度级联结构中的低成本高精度位置检测需求，支持双线性 Hall 信号生成、误差注入、补偿算法仿真、动态工况验证和结果导出。", size=11)
    add_heading(doc, "背景", 2)
    add_para(doc, "在手持相机云台、多自由度大角度级联系统及微型电机控制场景中，关节角度反馈对控制精度具有决定性影响。传统高精度编码器成本和结构约束较高，而双线性 Hall 传感器结合中空磁钢或环形磁钢，可在较低成本下实现连续角度检测。由于实际系统存在零偏、幅值不一致、安装角偏差、磁场谐波、ADC 量化和噪声等误差，必须配合信号补偿和动态估计算法才能达到工程可用的角度反馈精度。", size=11)
    add_heading(doc, "相关名词术语", 2)
    terms = [
        "线性 Hall 传感器：输出电压与外部磁场强度近似线性相关的磁传感器，可用于构造正余弦角度测量信号。",
        "Lissajous 相图：将两路归一化 Hall 信号作为横纵坐标绘制得到的轨迹。理想信号为单位圆，误差会导致偏心、椭圆、倾斜或波纹。",
        "ADC 量化：模拟电压经模数转换器转换为离散数字量的过程，会引入阶梯型量化误差。",
        "FFT：快速傅里叶变换，用于识别角度误差或信号中的主要谐波成分。",
        "α-β 滤波：基于角度和角速度的预测-校正估计方法，用于平滑角度测量并保持动态跟随。",
        "ESO：扩张状态观测器，可同时估计角度、速度和等效扰动，用于处理未建模动态误差。"
    ]
    for t in terms:
        add_para(doc, t, size=11)

    add_heading(doc, "参考资料", 2)
    refs = [
        "Position Extraction of Ultralow-Speed Gimbal Servo System With Linear Hall Sensors, IEEE Transactions on Industrial Electronics, 2021.",
        "TI Linear Hall Effect Sensor Angle Measurement 应用资料。",
        "本项目“双线性霍尔位置检测方案规划与学习路线”及相关 MATLAB/Simulink 仿真文件。"
    ]
    for r in refs:
        add_para(doc, r, size=11)

    add_heading(doc, "2 阅读指南", 1)
    add_heading(doc, "2.1 手册目的", 2)
    add_para(doc, "通过阅读本手册，用户可以了解软件的功能模块、参数含义、算法原理、仿真流程和结果导出方式，并可据此完成双线性 Hall 位置检测方案的演示、验证和二次开发。", size=11)
    add_heading(doc, "2.2 阅读对象", 2)
    add_para(doc, "本手册面向电机控制算法开发人员、云台传感器方案设计人员、软件著作权申请材料编写人员以及项目评审人员。", size=11)
    add_heading(doc, "2.3 运行环境", 2)
    add_para(doc, "硬件环境：普通 PC 机即可运行。", size=11)
    add_para(doc, "软件环境：推荐 Windows 10/11 操作系统，使用 Microsoft Edge、Google Chrome 或其他现代浏览器打开 HTML 文件。软件采用 HTML、CSS 和 JavaScript 编写，为离线单文件应用，不依赖网络和外部服务器。", size=11)

    add_heading(doc, "3 系统描述", 1)
    add_heading(doc, "3.1 软件开发方式", 2)
    add_para(doc, "本软件采用前端单文件架构实现，核心文件为 dual_hall_position_detection_app.html。软件使用 JavaScript 完成信号生成、误差注入、补偿算法、工况仿真和指标计算，使用 Canvas 绘制电压波形、相图、FFT、误差曲线、指标图和三维安装示意。", size=11)
    add_figure(doc, figs[0], "图1 软件功能结构图")

    add_heading(doc, "3.2 信号生成与误差注入", 2)
    add_para(doc, "软件根据 Hall 极对数、机械转速、安装角、仿真时长和采样频率生成磁场角与机械角，并按双线性 Hall 模型生成两路正余弦电压信号。用户可选择注入零偏、幅值不一致、相位不正交、高次谐波、白噪声和 ADC 量化误差。", size=11)
    add_para(doc, "基本信号模型如下：Hs = Os + As sin(θmag) + 谐波项 + ns；Hc = Oc + Ac cos(θmag + Δφ) + 谐波项 + nc。其中 Δφ 同时包含机械安装角偏离 90°造成的等效相位误差。", size=11)
    add_figure(doc, figs[2], "图2 中空磁钢与双线性 Hall 安装示意")

    add_heading(doc, "3.3 补偿算法", 2)
    add_para(doc, "软件提供确定性补偿和动态估计两类算法。确定性补偿包括零偏补偿、幅值归一化、相位补偿和谐波补偿，用于修正结构、安装和磁场周期畸变带来的可重复误差。动态估计包括一阶低通滤波、α-β 滤波和离散 ESO，用于处理随机噪声、采样误差、冲击和变速工况。", size=11)
    add_para(doc, "主要计算链路为：两路 Hall 信号 → 零偏/幅值归一化 → 相位/安装角补偿 → 谐波重构补偿 → atan2 解角 → α-β 或 ESO 动态估计 → 误差指标输出。", size=11)
    add_figure(doc, figs[1], "图3 信号处理与补偿流程图")

    add_heading(doc, "3.4 界面设计", 2)
    add_para(doc, "软件界面采用浅蓝色主题，左侧为参数配置区，右侧为结果显示区。用户可以通过输入框和滑块调节 Hall 极对数、安装角、幅值、零偏、相位误差、谐波强度、ADC 位数、噪声强度、采样频率、滤波参数和工况参数。", size=11)
    add_para(doc, "结果显示区包括动态三维安装示意、电压波形、Lissajous 相图、FFT 主要谐波、机械角解算误差、误差柱状图、雷达图、算法流程和原理介绍。", size=11)
    add_figure(doc, figs[3], "图4 软件界面布局示意")

    add_heading(doc, "3.5 工况仿真与结果导出", 2)
    add_para(doc, "软件支持恒速、变速、冲击和负载扰动工况。用户可通过工况参数调整速度倍率、冲击角度和负载降速比例，并观察不同补偿算法在动态场景下的误差变化。软件支持导出 PNG 汇总图、CSV 误差指标表和 JSON 仿真参数文件，便于报告撰写和结果复现。", size=11)

    add_heading(doc, "4 系统可行性分析", 1)
    add_para(doc, "从功能完整性看，软件覆盖了双线性 Hall 方案从传感器建模、误差注入、补偿算法到结果评价的完整链路，能够作为方案论证、算法对比和软著展示的基础软件。", size=11)
    add_para(doc, "从算法可行性看，软件将零偏、幅值、相位、谐波等确定性误差与噪声、量化、动态扰动等随机或工况误差分开处理，符合真实传感器标定和在线估计的工程流程。", size=11)
    add_para(doc, "从工程应用看，软件参数与项目中双线性 Hall、16 bit ADC、中空磁钢、多级关节角度反馈等需求相对应，可为后续实机采样、在线标定和与 IMU/控制器融合提供算法验证入口。", size=11)
    add_para(doc, "从扩展性看，软件采用模块化 JavaScript 实现，后续可继续加入实测数据导入、在线标定、温漂补偿、IMU 融合、控制器接口和批量报告生成等功能。", size=11)

    doc.save(str(OUT_DOCX))
    return OUT_DOCX


if __name__ == "__main__":
    out = build_doc()
    print(out)
