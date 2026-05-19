from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "线性霍尔传感器" / "线性霍尔高精度位置检测与补偿方法章节.md"
DST = ROOT / "线性霍尔传感器" / "线性霍尔高精度位置检测与补偿方法章节.docx"


def set_font(run, name="宋体", size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def clean_inline(text):
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "")
    return text


def add_code_block(doc, code, kind):
    label = {
        "math": "公式",
        "mermaid": "流程图（Mermaid 源码）",
    }.get(kind, "代码")
    p = doc.add_paragraph()
    r = p.add_run(label + "：")
    set_font(r, bold=True)

    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line if line else " ")
        set_font(r, name="Consolas", size=9)
        r.font.color.rgb = RGBColor(40, 40, 40)


def add_table(doc, table_lines):
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return

    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"

    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            value = clean_inline(row[j]) if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(value)
            set_font(r, size=9, bold=(i == 0))
            if i == 0:
                set_cell_shading(cell, "D9EAF7")


def convert():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip():
            i += 1
            continue

        if line.startswith("```"):
            kind = line.strip().strip("`").strip()
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            add_code_block(doc, "\n".join(block), kind)
            continue

        if line.lstrip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 4)
            p = doc.add_heading(clean_inline(heading.group(2)), level=level)
            for run in p.runs:
                set_font(run, name="黑体", bold=True)
            i += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(clean_inline(bullet.group(1)))
            set_font(r)
            i += 1
            continue

        numbered = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(clean_inline(numbered.group(2)))
            set_font(r)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(clean_inline(line))
        set_font(r)
        i += 1

    doc.save(DST)
    print(DST)


if __name__ == "__main__":
    convert()
