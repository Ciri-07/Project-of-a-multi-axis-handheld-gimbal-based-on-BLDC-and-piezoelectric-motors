from pathlib import Path
import re
import subprocess
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "线性霍尔传感器"
SRC = FOLDER / "线性霍尔高精度位置检测与补偿方法章节.md"
TMP = FOLDER / "_linear_hall_docx_build.md"
OUT = FOLDER / "线性霍尔高精度位置检测与补偿方法章节.docx"

FIGURES = [
    ("图 1  闭环自校准与数据驱动补偿流程图", "figures/linear_hall_method1_self_calibration.png"),
    ("图 2  归一化、谐波重构与 ESO 复合信号提取流程图", "figures/linear_hall_method2_sem_eso.png"),
    ("图 3  传感器与电机结构一体化自位置检测流程图", "figures/linear_hall_method3_structure.png"),
    ("图 4  线性 Hall 与 IMU 高精度检测总体融合框图", "figures/linear_hall_overall_fusion.png"),
]


def find_pandoc():
    for p in os.environ.get("PATH", "").split(os.pathsep):
        exe = Path(p) / "pandoc.exe"
        if exe.exists():
            return exe
    root = Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "WinGet" / "Packages"
    if root.exists():
        hits = list(root.rglob("pandoc.exe"))
        if hits:
            return hits[0]
    raise RuntimeError("pandoc.exe not found")


def replace_mermaid_blocks(text):
    idx = 0

    def repl(match):
        nonlocal idx
        if idx >= len(FIGURES):
            return match.group(0)
        caption, img = FIGURES[idx]
        idx += 1
        return f"![{caption}]({img})\n\n*{caption}*"

    return re.sub(r"```mermaid\s+.*?```", repl, text, flags=re.S)


def set_run_font(run, name="宋体", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    """Set individual cell borders. Use val='nil' to remove a border."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        spec = kwargs.get(edge)
        if spec is None:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in spec.items():
            element.set(qn(f"w:{key}"), str(value))


def apply_three_line_table_style(table):
    """Academic three-line table: top rule, mid rule, bottom rule; no verticals."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    try:
        table.style = "Table Normal"
    except KeyError:
        pass

    row_count = len(table.rows)
    nil = {"val": "nil"}
    top_rule = {"val": "single", "sz": "14", "color": "000000", "space": "0"}
    mid_rule = {"val": "single", "sz": "9", "color": "000000", "space": "0"}
    bottom_rule = {"val": "single", "sz": "14", "color": "000000", "space": "0"}

    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell, top=90, start=120, bottom=90, end=120)
            set_cell_shading(cell, "FFFFFF")
            set_cell_border(cell, top=nil, start=nil, bottom=nil, end=nil)

            if ri == 0:
                set_cell_border(cell, top=top_rule, bottom=mid_rule)
            if ri == row_count - 1:
                set_cell_border(cell, bottom=bottom_rule)

            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    if ri == 0:
                        set_run_font(run, "黑体", 9.5, True, "000000")
                    else:
                        set_run_font(run, "宋体", 9, False, "111827")


def style_document(path):
    doc = Document(path)

    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        if style_name in doc.styles:
            st = doc.styles[style_name]
            st.font.name = "黑体"
            st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    for p in doc.paragraphs:
        if p.style.name == "Caption" or p.text.startswith("图 "):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, "宋体", 9, color="475569")
        elif p.style.name == "Normal":
            p.paragraph_format.line_spacing = 1.18
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                set_run_font(run, "宋体", 10.5)

    for table in doc.tables:
        apply_three_line_table_style(table)

    for shape in doc.inline_shapes:
        max_width = Cm(16.2)
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)

    doc.save(path)


def main():
    text = SRC.read_text(encoding="utf-8")
    text = replace_mermaid_blocks(text)
    TMP.write_text(text, encoding="utf-8")

    pandoc = find_pandoc()
    subprocess.run(
        [
            str(pandoc),
            TMP.name,
            "-o",
            OUT.name,
            "--from",
            "markdown+tex_math_dollars",
            "--to",
            "docx",
        ],
        cwd=FOLDER,
        check=True,
    )
    style_document(OUT)
    TMP.unlink(missing_ok=True)
    print(OUT)


if __name__ == "__main__":
    main()
