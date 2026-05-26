from pathlib import Path
from copy import deepcopy

from docx import Document


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "_template_source_code.docx"
SOURCE = ROOT / "_dual_hall_position_detection_app.html"
OUTPUT = ROOT / "_dual_hall_source_code.docx"


def clear_document_body(document):
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def copy_paragraph_format(dst, src):
    dst.style = src.style
    dst.alignment = src.alignment

    src_format = src.paragraph_format
    dst_format = dst.paragraph_format
    fields = [
        "left_indent",
        "right_indent",
        "first_line_indent",
        "space_before",
        "space_after",
        "line_spacing",
        "line_spacing_rule",
        "keep_together",
        "keep_with_next",
        "page_break_before",
        "widow_control",
    ]
    for field in fields:
        try:
            setattr(dst_format, field, getattr(src_format, field))
        except Exception:
            pass


def copy_run_format(dst, src):
    if src is None:
        return
    if src._r.rPr is not None:
        dst._r.get_or_add_rPr()
        for child in list(dst._r.rPr):
            dst._r.rPr.remove(child)
        for child in list(src._r.rPr):
            dst._r.rPr.append(deepcopy(child))


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Missing template: {TEMPLATE}")
    if not SOURCE.exists():
        raise FileNotFoundError(f"Missing source file: {SOURCE}")

    document = Document(str(TEMPLATE))
    source_text = SOURCE.read_text(encoding="utf-8")
    source_lines = source_text.splitlines()

    proto_paragraph = document.paragraphs[0] if document.paragraphs else None
    proto_run = proto_paragraph.runs[0] if proto_paragraph and proto_paragraph.runs else None

    clear_document_body(document)

    for line in source_lines:
        paragraph = document.add_paragraph()
        if proto_paragraph is not None:
            copy_paragraph_format(paragraph, proto_paragraph)
        run = paragraph.add_run(line)
        copy_run_format(run, proto_run)

    document.save(str(OUTPUT))
    print(f"written: {OUTPUT}")
    print(f"source lines: {len(source_lines)}")


if __name__ == "__main__":
    main()
